import pandas as pd
import os
import openpyxl
from datetime import datetime

def salvar_dados_detalhados(lista_dados, pasta_destino, cnpj):
    """
    Recebe uma lista de dicionários com os dados da tabela e salva em um arquivo Excel.
    Se o arquivo já existir, anexa os novos dados.
    """
    if not lista_dados:
        print("⚠️ Nenhun dado para salvar.")
        return False

    try:
        # Nome do arquivo
        data_atual = datetime.now().strftime("%Y-%m-%d")
        nome_arquivo = f"extracao_fgts_digital_{cnpj}_{data_atual}.xlsx"
        caminho_arquivo = os.path.join(pasta_destino, nome_arquivo)

        # Converte lista de dicionários em DataFrame
        df_novo = pd.DataFrame(lista_dados)

        # Adiciona metadados
        df_novo['Data Extracao'] = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        df_novo['CNPJ'] = cnpj

        # Verifica se o arquivo já existe para anexar ou criar novo
        if os.path.exists(caminho_arquivo):
            try:
                df_antigo = pd.read_excel(caminho_arquivo)
                df_final = pd.concat([df_antigo, df_novo], ignore_index=True)
            except Exception as e:
                print(f"⚠️ Erro ao ler arquivo existente, criando um novo: {e}")
                df_final = df_novo
        else:
            df_final = df_novo

        # Salva o DataFrame em Excel
        df_final.to_excel(caminho_arquivo, index=False, engine='openpyxl')
        print(f"✅ Dados salvos com sucesso em: {caminho_arquivo}")
        return True

    except Exception as e:
        print(f"❌ Erro ao salvar planilha Excel: {e}")
        return False

def atualizar_valor_guia_original(caminho_xlsx, cnpj, mapa_valores):
    """
    Localiza o CNPJ na planilha original e atualiza as colunas de valores conforme a extração.
    mapa_valores: dicionário { "MM/AAAA": valor, "13º/AAAA": valor }
    """
    if not caminho_xlsx or not os.path.exists(caminho_xlsx):
        print(f"⚠️ Arquivo original não encontrado: {caminho_xlsx}")
        return False

    try:
        wb = openpyxl.load_workbook(caminho_xlsx)
        sheet = wb.active

        # Encontrar cabeçalhos (limpando espaços e forçando maiúsculas)
        headers = {}
        for cell in sheet[1]:
            if cell.value:
                headers[str(cell.value).strip().upper()] = cell.column

        col_cnpj = headers.get("CNPJ")
        col_valor_12 = headers.get("VALOR_GUIA_FGTS")
        col_valor_13 = headers.get("VALOR_GUIA_FGTS13")
        col_total_guia = headers.get("TOTAL_GUIA")
        col_total_original = headers.get("TOTAL") # Coluna com valor esperado
        col_validacao = headers.get("VALIDACAO")

        if col_cnpj is None:
            print("❌ Coluna 'CNPJ' não encontrada na planilha original.")
            return False

        print(f"📊 Colunas encontradas: CNPJ={col_cnpj}, VALOR_12={col_valor_12}, VALOR_13={col_valor_13}, TOTAL_GUIA={col_total_guia}, VALIDACAO={col_validacao}")

        cnpj_limpo = ''.join(filter(str.isdigit, str(cnpj)))

        # Função auxiliar interna para limpar valor monetário
        def limpar_valor(v):
            if v is None: return 0.0
            if isinstance(v, (int, float)): return float(v)
            s = str(v).replace('R$', '').replace(' ', '').replace('.', '').replace(',', '.')
            try: return float(s)
            except: return 0.0

        encontrado = False
        for row in range(2, sheet.max_row + 1):
            valor_celula_cnpj = str(sheet.cell(row=row, column=col_cnpj).value)
            cnpj_celula_limpo = ''.join(filter(str.isdigit, valor_celula_cnpj))

            if cnpj_celula_limpo == cnpj_limpo:
                # Extrai valores do mapa
                valor_mensal = 0.0
                valor_13 = 0.0
                total_extraido = 0.0
                
                # Se o mapa contiver o total consolidado (capturado via Regex específica de Total/Total da Guia)
                if "TOTAL_CONSOLIDADO" in mapa_valores:
                    total_extraido = mapa_valores["TOTAL_CONSOLIDADO"]
                
                for comp, val in mapa_valores.items():
                    if comp == "TOTAL_CONSOLIDADO": continue
                    if "13º" in comp:
                        valor_13 = val
                    else:
                        valor_mensal = val
                
                # Se não houver total consolidado, faz o fallback para soma (padrão antigo)
                if total_extraido == 0:
                    total_extraido = valor_mensal + valor_13

                print(f"📝 Atualizando CNPJ {cnpj}: Mensal={valor_mensal}, 13º={valor_13}, Total={total_extraido}")

                # Atualiza as colunas de extração
                if col_valor_12: sheet.cell(row=row, column=col_valor_12).value = valor_mensal
                if col_valor_13: sheet.cell(row=row, column=col_valor_13).value = valor_13
                if col_total_guia: sheet.cell(row=row, column=col_total_guia).value = total_extraido
                
                # NOVO: Grava data de pagamento se existir no mapa e a coluna estiver presente
                if "DATA_PAGAMENTO" in mapa_valores:
                    col_data_pgto = headers.get("DATA_PAGAMENTO") or headers.get("VENCIMENTO")
                    if col_data_pgto:
                        sheet.cell(row=row, column=col_data_pgto).value = mapa_valores["DATA_PAGAMENTO"]
                        print(f"📅 Data de pagamento ({mapa_valores['DATA_PAGAMENTO']}) gravada na coluna {col_data_pgto}")

                # Validação
                if col_total_original and col_validacao:
                    try:
                        v_esperado = limpar_valor(sheet.cell(row=row, column=col_total_original).value)
                        print(f"⚖️ Comparando Esperado({v_esperado}) com Extraído({total_extraido})")
                        
                        if abs(v_esperado - total_extraido) < 0.02: # Tolerância de 2 cents
                            sheet.cell(row=row, column=col_validacao).value = "SIM"
                            print(f"✅ Validação OK para CNPJ {cnpj}.")
                        else:
                            sheet.cell(row=row, column=col_validacao).value = "DIVERGENTE"
                            print(f"⚠️ Valores divergentes para CNPJ {cnpj}")
                    except Exception as ve:
                        print(f"⚠️ Erro ao validar: {ve}")

                encontrado = True
                break

        if not encontrado:
            print(f"⚠️ CNPJ {cnpj} não encontrado.")
            return False

        wb.save(caminho_xlsx)
        print(f"💾 Planilha salva com sucesso.")
        
        # Retorna sucesso se validado
        if col_validacao:
            validacao_final = str(sheet.cell(row=row, column=col_validacao).value)
            return validacao_final == "SIM"
        return True

    except Exception as e:
        print(f"❌ Erro ao atualizar planilha original: {e}")
        return False

def marcar_guia_nao_encontrada(caminho_xlsx, cnpj):
    """
    Localiza o CNPJ na planilha original e define a coluna VALIDACAO como 'Não há guia para emitir'.
    """
    if not caminho_xlsx or not os.path.exists(caminho_xlsx):
        print(f"⚠️ Arquivo original não encontrado: {caminho_xlsx}")
        return False

    try:
        wb = openpyxl.load_workbook(caminho_xlsx)
        sheet = wb.active

        headers = {}
        for cell in sheet[1]:
            if cell.value:
                headers[str(cell.value).strip().upper()] = cell.column

        col_cnpj = headers.get("CNPJ")
        col_validacao = headers.get("VALIDACAO")

        if col_cnpj is None or col_validacao is None:
            print("❌ Colunas necessárias não encontradas na planilha para marcar erro.")
            return False

        cnpj_limpo = ''.join(filter(str.isdigit, str(cnpj)))
        encontrado = False

        for row in range(2, sheet.max_row + 1):
            valor_celula_cnpj = str(sheet.cell(row=row, column=col_cnpj).value)
            cnpj_celula_limpo = ''.join(filter(str.isdigit, valor_celula_cnpj))

            if cnpj_celula_limpo == cnpj_limpo:
                sheet.cell(row=row, column=col_validacao).value = "Não há guia para emitir"
                encontrado = True
                break

        if encontrado:
            wb.save(caminho_xlsx)
            print(f"✅ Planilha atualizada para CNPJ {cnpj}: Não há guia para emitir.")
            return True
        else:
            print(f"⚠️ CNPJ {cnpj} não encontrado para atualizar status.")
            return False

    except Exception as e:
        print(f"❌ Erro ao atualizar planilha (guia não encontrada): {e}")
        return False

def marcar_status_validacao(caminho_xlsx, cnpj, mensagem):
    """
    Localiza o CNPJ na planilha original e define a coluna VALIDACAO com a mensagem fornecida.
    """
    if not caminho_xlsx or not os.path.exists(caminho_xlsx):
        print(f"⚠️ Arquivo original não encontrado: {caminho_xlsx}")
        return False

    try:
        wb = openpyxl.load_workbook(caminho_xlsx)
        sheet = wb.active

        headers = {}
        for cell in sheet[1]:
            if cell.value:
                headers[str(cell.value).strip().upper()] = cell.column

        col_cnpj = headers.get("CNPJ")
        col_validacao = headers.get("VALIDACAO")

        if col_cnpj is None or col_validacao is None:
            print("❌ Colunas necessárias não encontradas na planilha para marcar status.")
            return False

        cnpj_limpo = ''.join(filter(str.isdigit, str(cnpj)))
        encontrado = False

        for row in range(2, sheet.max_row + 1):
            valor_celula_cnpj = str(sheet.cell(row=row, column=col_cnpj).value)
            cnpj_celula_limpo = ''.join(filter(str.isdigit, valor_celula_cnpj))

            if cnpj_celula_limpo == cnpj_limpo:
                sheet.cell(row=row, column=col_validacao).value = mensagem
                encontrado = True
                break

        if encontrado:
            wb.save(caminho_xlsx)
            print(f"✅ Planilha atualizada para CNPJ {cnpj}: {mensagem}")
            return True
        else:
            print(f"⚠️ CNPJ {cnpj} não encontrado para atualizar status.")
            return False

    except Exception as e:
        print(f"❌ Erro ao marcar status de validação: {e}")
        return False
